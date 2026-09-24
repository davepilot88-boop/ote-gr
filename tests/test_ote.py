import importlib.util
import json
import unittest
from datetime import date
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
spec=importlib.util.spec_from_file_location('ote',ROOT/'.github/workflows/update_ote.py')
ote=importlib.util.module_from_spec(spec)
spec.loader.exec_module(ote)
FIXTURES=ROOT/'tests/fixtures'

def fixture(day):
    return (FIXTURES/f'DT_15MIN_{day:%d_%m_%Y}_CZ.xlsx').read_bytes()

class ParserTests(unittest.TestCase):
    def test_real_exports_and_dst(self):
        for day,count in [(date(2026,3,29),23),(date(2026,3,30),24),(date(2025,10,26),25)]:
            with self.subTest(day=day):
                p=ote.parse_export(fixture(day),day)
                self.assertEqual(len(p['prices']),count)
                self.assertEqual(len(p['intervals']),count)
        p=ote.parse_export(fixture(date(2026,3,29)),date(2026,3,29))
        self.assertEqual(p['prices'][12:15],[39.56,18.10,19.23])
        p=ote.parse_export(fixture(date(2025,10,26)),date(2025,10,26))
        self.assertIn('02a',p['intervals'][2]['label'])
        self.assertIn('02b',p['intervals'][3]['label'])

    def mutated(self, edits):
        book=load_workbook(BytesIO(fixture(date(2026,3,30))))
        for address,value in edits.items():book.active[address]=value
        buf=BytesIO();book.save(buf);book.close();return buf.getvalue()

    def test_rejects_schema_date_missing_nonfinite_and_intervals(self):
        for edits in [{'A20':'Výsledky denního trhu ČR - 29.03.2026'}, {'L22':'Množství'},
                      {'L24':None},{'L24':'NaN'},{'L24':False},{'B25':'00:00-00:15'},
                      {'A25':1},{'L25':99}, {'C24':None}]:
            with self.subTest(edits=edits), self.assertRaises(ote.DataError):
                ote.parse_export(self.mutated(edits),date(2026,3,30))

    def test_negative_zero_and_equal_neighboring_hours(self):
        p=ote.parse_export(self.mutated({f'{col}{i}':-12.5 for i in range(24,32) for col in ['C','L']}),date(2026,3,30))
        self.assertEqual(p['prices'][:2],[-12.5,-12.5])
        p=ote.parse_export(self.mutated({f'{col}{i}':0 for i in range(24,28) for col in ['C','L']}),date(2026,3,30))
        self.assertEqual(p['prices'][0],0)

    def test_sanity_guard(self):
        day=date(2026,3,29)
        book=load_workbook(BytesIO(fixture(day)))
        for i in range(72,76):
            book.active[f'L{i}']=40
            book.active[f'C{i}']=40
        buf=BytesIO();book.save(buf);book.close()
        with self.assertRaisesRegex(ote.DataError,'Sanity'):ote.parse_export(buf.getvalue(),day)

    def test_unpublished_requires_explicit_marker_and_date(self):
        day=date(2026,3,30)
        h='Výsledky denního trhu ČR - 30.03.2026 '
        self.assertIsNone(ote.export_link(h+'Pro zvolený filtr nejsou dostupná data',day))
        for html in [h,'<h1>Server error</h1>',h+'<a href="wrong/DT_15MIN_file.xlsx">x</a>']:
            with self.assertRaises(ote.DataError):ote.export_link(html,day)

    def test_error_preserves_previous_file(self):
        with TemporaryDirectory() as temp:
            out=Path(temp)
            for name in ['today','tomorrow']:(out/f'{name}.json').write_text('previous valid bytes')
            def fail(day):raise OSError('OTE unavailable')
            self.assertEqual(ote.update(out,date(2026,3,30),fail),1)
            for name in ['today','tomorrow']:self.assertEqual((out/f'{name}.json').read_text(),'previous valid bytes')
            self.assertEqual(json.loads((out/'status.json').read_text())['days']['tomorrow']['state'],'error')

    def test_unpublished_preserves_old_date_and_saves_status(self):
        day=date(2026,3,30)
        with TemporaryDirectory() as temp:
            out=Path(temp);(out/'tomorrow.json').write_text('old dated payload')
            p=ote.parse_export(fixture(day),day)
            self.assertEqual(ote.update(out,day,lambda d:p if d==day else None),0)
            self.assertEqual((out/'tomorrow.json').read_text(),'old dated payload')
            state=json.loads((out/'status.json').read_text())
            self.assertEqual(state['days']['tomorrow'],{'date':'31.03.2026','state':'unpublished'})

if __name__=='__main__':unittest.main()
